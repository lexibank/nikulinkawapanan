import pathlib
from openpyxl import load_workbook
from clldutils.misc import slug
from pylexibank import Dataset as BaseDataset
from pylexibank import progressbar as pb
from lingpy import Wordlist


class Dataset(BaseDataset):
    dir = pathlib.Path(__file__).parent
    id = "nikulinkawapanan"

    writer_options = {'keep_languages': False, 'keep_parameters': False}

    def cmd_download(self, _):
        ...

    def cmd_makecldf(self, args):
        # add bib
        args.writer.add_sources()
        args.log.info("added sources")

        # # add concept
        # concepts = {}
        # for concept in self.conceptlists[0].concepts.values():
        #     idx = concept.id.split("-")[-1] + "_" + slug(concept.english)
        #     args.writer.add_concept(
        #         ID=idx,
        #         Name=concept.english,
        #         Concepticon_ID=concept.concepticon_id,
        #         Concepticon_Gloss=concept.concepticon_gloss,
        #     )
        #     concepts[concept.english] = idx

        # args.log.info("added concepts")

        # # add languages
        # languages = {}
        # for language in self.languages:
        #     args.writer.add_language(
        #             ID=language["ID"],
        #             Name=language["Name"],
        #             Glottocode=language["Glottocode"],
        #             SubGroup=language["SubGroup"]
        #             )
        #     languages[language["ID"]] = language["Name"]

        # args.log.info("added languages")


        workbook = load_workbook(self.raw_dir.joinpath("Nikulin-2022-Kawapanan.xlsx"))
        sheet = workbook.active

        header = [cell.value for cell in sheet[1]]

        set_col = header.index("set")
        entry_col = header.index("entry")
        gloss_col = header.index("gloss")

        # everything between entry and gloss are doculects
        doculect_cols = {
            header[i]: i
            for i in range(entry_col + 1, gloss_col)
        }

        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            set_id = row[set_col].value
            entry = row[entry_col].value
            gloss = row[gloss_col].value

            for doculect, col_idx in doculect_cols.items():
                value = row[col_idx].value

                print(
                    f"Set: {set_id}, "
                    f"Entry: {entry}, "
                    f"Doculect: {doculect}, "
                    f"Value: {value}, "
                    f"Gloss: {gloss}"
                )

                # if value:
                #     args.writer.add_forms_from_value(
                #         Language_ID=languages[doculect],
                #         Parameter_ID=concepts[concept],
                #         Value=value,
                #         Source=sources[doculect]
                #         )
