from openpyxl import load_workbook
import csv

workbook = load_workbook("Nikulin-2022-Kawapanan.xlsx")
sheet = workbook.active

header = [cell.value for cell in sheet[1]]

set_col = header.index("set")
entry_col = header.index("entry")
gloss_col = header.index("gloss")

# everything between entry and gloss are doculects
doculect_cols = {
    header[i]: i
    for i in range(entry_col + 1, gloss_col)
}

data = [['set', 'entry', 'doculect', 'value', 'concept']]
concepts = [['gloss']]
visited = []
for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
    set_id = row[set_col].value
    entry = row[entry_col].value
    gloss = row[gloss_col].value

    for doculect, col_idx in doculect_cols.items():
        value = row[col_idx].value

        if [doculect, value, gloss] not in visited:
            visited.append([doculect, value, gloss])
            data.append([set_id, entry, doculect, value, gloss])
            print(
                f"Set: {set_id}, "
                f"Entry: {entry}, "
                f"Doculect: {doculect}, "
                f"Value: {value}, "
                f"Gloss: {gloss}"
            )

        if [gloss] not in concepts:
            concepts.append([gloss])

with open("raw.tsv", 'w', encoding="utf8") as f:
    writer = csv.writer(f, delimiter='\t')
    writer.writerows(data)

with open("concepts.tsv", 'w', encoding="utf8") as f:
    writer = csv.writer(f, delimiter='\t')
    writer.writerows(concepts)
